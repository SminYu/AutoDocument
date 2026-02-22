from __future__ import annotations

import argparse
import json
import logging
import os
import re
import shutil
import uuid
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pytesseract

from .utils import business_checksum, normalize_date, to_int_amount


LOGGER = logging.getLogger("autodocument")


@dataclass
class OCRLine:
    text: str
    confidence: float
    bbox: tuple[int, int, int, int]


@dataclass
class ExtractedFields:
    merchant_name: str | None = None
    business_number: str | None = None
    approved_at: str | None = None
    supply_amount: int | None = None
    vat_amount: int | None = None
    total_amount: int | None = None
    approval_no: str | None = None
    reasons: dict[str, str] = field(default_factory=dict)


@dataclass
class ValidationResult:
    status: str
    checks: list[dict[str, Any]]


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha256(path: Path) -> str:
    h = sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _order_points(pts: np.ndarray) -> np.ndarray:
    rect = np.zeros((4, 2), dtype="float32")
    s = pts.sum(axis=1)
    rect[0] = pts[np.argmin(s)]
    rect[2] = pts[np.argmax(s)]

    diff = np.diff(pts, axis=1)
    rect[1] = pts[np.argmin(diff)]
    rect[3] = pts[np.argmax(diff)]
    return rect


def _four_point_transform(image: np.ndarray, pts: np.ndarray) -> np.ndarray:
    rect = _order_points(pts)
    (tl, tr, br, bl) = rect

    width_a = np.linalg.norm(br - bl)
    width_b = np.linalg.norm(tr - tl)
    max_w = max(int(width_a), int(width_b))

    height_a = np.linalg.norm(tr - br)
    height_b = np.linalg.norm(tl - bl)
    max_h = max(int(height_a), int(height_b))

    dst = np.array(
        [[0, 0], [max_w - 1, 0], [max_w - 1, max_h - 1], [0, max_h - 1]],
        dtype="float32",
    )

    mat = cv2.getPerspectiveTransform(rect, dst)
    return cv2.warpPerspective(image, mat, (max_w, max_h))


def preprocess_image(input_path: Path, out_dir: Path) -> list[Path]:
    image = cv2.imread(str(input_path))
    if image is None:
        raise ValueError(f"Cannot read input image: {input_path}")

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blurred = cv2.GaussianBlur(gray, (5, 5), 0)
    edged = cv2.Canny(blurred, 50, 200)

    contours, _ = cv2.findContours(edged.copy(), cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    contours = sorted(contours, key=cv2.contourArea, reverse=True)[:5]

    warped = image.copy()
    for contour in contours:
        peri = cv2.arcLength(contour, True)
        approx = cv2.approxPolyDP(contour, 0.02 * peri, True)
        if len(approx) == 4:
            warped = _four_point_transform(image, approx.reshape(4, 2))
            break

    gray_warped = cv2.cvtColor(warped, cv2.COLOR_BGR2GRAY)
    coords = np.column_stack(np.where(gray_warped > 0))
    angle = cv2.minAreaRect(coords)[-1] if coords.size else 0
    angle = -(90 + angle) if angle < -45 else -angle
    h, w = gray_warped.shape[:2]
    m = cv2.getRotationMatrix2D((w // 2, h // 2), angle, 1.0)
    deskewed = cv2.warpAffine(gray_warped, m, (w, h), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)

    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
    enhanced = clahe.apply(deskewed)
    denoised = cv2.medianBlur(enhanced, 3)
    bw = cv2.adaptiveThreshold(
        denoised, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 11
    )
    upscaled = cv2.resize(bw, None, fx=2.0, fy=2.0, interpolation=cv2.INTER_CUBIC)

    out_dir.mkdir(parents=True, exist_ok=True)
    variants = {
        "01_warped.jpg": cv2.cvtColor(gray_warped, cv2.COLOR_GRAY2BGR),
        "02_deskewed.jpg": cv2.cvtColor(deskewed, cv2.COLOR_GRAY2BGR),
        "03_enhanced.jpg": cv2.cvtColor(denoised, cv2.COLOR_GRAY2BGR),
        "04_bw.jpg": cv2.cvtColor(bw, cv2.COLOR_GRAY2BGR),
        "05_upscaled.jpg": cv2.cvtColor(upscaled, cv2.COLOR_GRAY2BGR),
    }
    paths: list[Path] = []
    for name, img in variants.items():
        p = out_dir / name
        cv2.imwrite(str(p), img)
        paths.append(p)
    return paths


def run_ocr(images: list[Path], lang: str = "kor+eng") -> tuple[list[OCRLine], str]:
    lines: list[OCRLine] = []
    text_blobs: list[str] = []

    for image_path in images:
        img = cv2.imread(str(image_path))
        data = pytesseract.image_to_data(img, lang=lang, output_type=pytesseract.Output.DICT, config="--oem 1 --psm 6")
        n = len(data["text"])
        for idx in range(n):
            text = data["text"][idx].strip()
            if not text:
                continue
            conf = float(data["conf"][idx]) if data["conf"][idx] != "-1" else 0.0
            bbox = (
                int(data["left"][idx]),
                int(data["top"][idx]),
                int(data["width"][idx]),
                int(data["height"][idx]),
            )
            lines.append(OCRLine(text=text, confidence=conf, bbox=bbox))
        text_blobs.append(pytesseract.image_to_string(img, lang=lang, config="--oem 1 --psm 6"))

    merged_text = "\n".join(text_blobs)
    return lines, merged_text


def extract_fields(ocr_text: str) -> ExtractedFields:
    field = ExtractedFields()

    bn = re.search(r"(\d{3}-\d{2}-\d{5}|\d{10})", ocr_text)
    if bn:
        digits = re.sub(r"[^0-9]", "", bn.group(1))
        field.business_number = f"{digits[:3]}-{digits[3:5]}-{digits[5:]}"

    date_match = re.search(r"(20\d{2}[./-]\d{1,2}[./-]\d{1,2}|20\d{6})", ocr_text)
    if date_match:
        field.approved_at = normalize_date(date_match.group(1))

    total = re.search(r"(?:합계|총액|Total)\s*[: ]?\s*([\d,]+)", ocr_text, re.IGNORECASE)
    supply = re.search(r"(?:공급가액|공급가)\s*[: ]?\s*([\d,]+)", ocr_text, re.IGNORECASE)
    vat = re.search(r"(?:부가세|VAT)\s*[: ]?\s*([\d,]+)", ocr_text, re.IGNORECASE)

    if total:
        field.total_amount = to_int_amount(total.group(1))
    if supply:
        field.supply_amount = to_int_amount(supply.group(1))
    if vat:
        field.vat_amount = to_int_amount(vat.group(1))

    app_no = re.search(r"(?:승인번호|Approval\s*No)\s*[: ]?\s*([A-Z0-9-]{6,})", ocr_text, re.IGNORECASE)
    if app_no:
        field.approval_no = app_no.group(1)

    first_line = next((line.strip() for line in ocr_text.splitlines() if line.strip()), None)
    field.merchant_name = first_line

    for key in [
        "merchant_name",
        "business_number",
        "approved_at",
        "supply_amount",
        "vat_amount",
        "total_amount",
        "approval_no",
    ]:
        if getattr(field, key) is None:
            field.reasons[key] = "not_found_by_rules"

    return field


def validate(fields: ExtractedFields) -> ValidationResult:
    checks: list[dict[str, Any]] = []

    formula_ok = (
        fields.total_amount is not None
        and fields.supply_amount is not None
        and fields.vat_amount is not None
        and fields.total_amount == fields.supply_amount + fields.vat_amount
    )
    checks.append({"name": "total_formula", "result": formula_ok})

    bn_ok = fields.business_number is not None and business_checksum(fields.business_number)
    checks.append({"name": "business_number_checksum", "result": bn_ok})

    if all(c["result"] for c in checks):
        status = "PASS"
    elif any(c["result"] for c in checks):
        status = "REVIEW"
    else:
        status = "FAIL"
    return ValidationResult(status=status, checks=checks)


def _replace_placeholders(text: str, fields: ExtractedFields) -> str:
    replacements = {
        "{{merchant_name}}": fields.merchant_name or "",
        "{{business_number}}": fields.business_number or "",
        "{{approved_at}}": fields.approved_at or "",
        "{{supply_amount}}": str(fields.supply_amount or ""),
        "{{vat_amount}}": str(fields.vat_amount or ""),
        "{{total_amount}}": str(fields.total_amount or ""),
        "{{approval_no}}": fields.approval_no or "",
    }
    for key, value in replacements.items():
        text = text.replace(key, value)
    return text


def fill_template(template_path: Path | None, output_dir: Path, fields: ExtractedFields) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    if template_path is None:
        out = output_dir / "filled_evidence.txt"
        out.write_text(json.dumps(asdict(fields), ensure_ascii=False, indent=2), encoding="utf-8")
        return out

    ext = template_path.suffix.lower()

    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            out = output_dir / "filled_evidence.txt"
            out.write_text(
                _replace_placeholders(template_path.read_text(encoding="utf-8"), fields),
                encoding="utf-8",
            )
            return out
        doc = Document(str(template_path))
        for paragraph in doc.paragraphs:
            paragraph.text = _replace_placeholders(paragraph.text, fields)
        out = output_dir / f"filled_{template_path.name}"
        doc.save(str(out))
        return out

    if ext == ".hwpx":
        out = output_dir / f"filled_{template_path.name}"
        with zipfile.ZipFile(template_path, "r") as zin, zipfile.ZipFile(out, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith(".xml"):
                    decoded = data.decode("utf-8", errors="ignore")
                    decoded = _replace_placeholders(decoded, fields)
                    data = decoded.encode("utf-8")
                zout.writestr(item, data)
        return out

    out = output_dir / f"filled_{template_path.stem}.txt"
    raw = template_path.read_text(encoding="utf-8")
    out.write_text(_replace_placeholders(raw, fields), encoding="utf-8")
    return out


def run_pipeline(input_path: Path, output_root: Path, template_path: Path | None = None) -> dict[str, Any]:
    document_id = str(uuid.uuid4())
    run_dir = output_root / document_id
    pre_dir = run_dir / "preprocessed"

    run_dir.mkdir(parents=True, exist_ok=True)

    raw_copy = run_dir / input_path.name
    shutil.copy2(input_path, raw_copy)

    preprocessed_images = preprocess_image(raw_copy, pre_dir)
    ocr_lines, ocr_text = run_ocr(preprocessed_images)
    fields = extract_fields(ocr_text)
    validation = validate(fields)
    evidence_path = fill_template(template_path, run_dir, fields)

    result = {
        "document_id": document_id,
        "source": {
            "file_name": input_path.name,
            "sha256": _sha256(raw_copy),
            "uploaded_at": _utc_now(),
        },
        "ocr": {
            "line_count": len(ocr_lines),
            "avg_confidence": round(
                sum(line.confidence for line in ocr_lines) / len(ocr_lines), 2
            )
            if ocr_lines
            else 0.0,
        },
        "extracted": asdict(fields),
        "validation": asdict(validation),
        "outputs": {
            "evidence_file": str(evidence_path),
            "preprocessed_images": [str(x) for x in preprocessed_images],
        },
        "audit": {
            "ocr_engine": "pytesseract",
            "extractor_version": "rule-v1",
            "rule_set_version": "2026.02",
        },
    }

    (run_dir / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "ocr_raw.txt").write_text(ocr_text, encoding="utf-8")

    LOGGER.info("Pipeline complete. Result: %s", run_dir / "result.json")
    return result


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="AutoDocument receipt pipeline")
    parser.add_argument("--input", required=True, type=Path, help="Input image path")
    parser.add_argument("--output", required=True, type=Path, help="Output directory")
    parser.add_argument("--template", type=Path, default=None, help="Template path (DOCX/HWPX/TXT)")
    return parser


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    args = build_arg_parser().parse_args()
    run_pipeline(args.input, args.output, args.template)


if __name__ == "__main__":
    main()
